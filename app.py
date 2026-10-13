import io

import re

from datetime import datetime
from zoneinfo import ZoneInfo

import openpyxl

from openpyxl.formatting.rule import CellIsRule

from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment

import pandas as pd

import streamlit as st



# Librerías para generación de PDF

from reportlab.lib import colors

from reportlab.lib.pagesizes import A4, landscape

from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet

from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle



# ---------------------------------------------------------

# CONFIGURACIÓN DE PÁGINA

# ---------------------------------------------------------

st.set_page_config(

    page_title="Seguimiento de Inversiones usando datos del SIAF",

    page_icon="🏛️",

    layout="wide",

)


# ---------------------------------------------------------
# ADAPTACIÓN VISUAL RESPONSIVE (PC + CELULAR)
# No modifica la lógica, cálculos ni generación de archivos.
# ---------------------------------------------------------
st.markdown(
    """
    <style>
    /* Contenedor general */
    .block-container {
        max-width: 1500px;
        padding-top: 1.4rem;
        padding-bottom: 2rem;
    }

    /* Tarjetas KPI: alto contraste en modo claro y oscuro */
    [data-testid="stMetric"] {
        background: #F1F5F9;
        border: 1px solid #CBD5E1;
        border-radius: 12px;
        padding: 0.8rem 1rem;
        min-height: 105px;
        box-shadow: 0 1px 2px rgba(15, 23, 42, 0.06);
    }

    [data-testid="stMetricLabel"],
    [data-testid="stMetricLabel"] p {
        color: #334155 !important;
        font-weight: 600 !important;
    }

    [data-testid="stMetricValue"],
    [data-testid="stMetricValue"] > div {
        color: #0F172A !important;
    }

    [data-testid="stMetricDelta"] {
        font-weight: 600 !important;
    }

    /* Botones cómodos en pantallas táctiles */
    .stDownloadButton > button,
    .stButton > button {
        min-height: 46px;
        border-radius: 10px;
        font-weight: 600;
    }

    /* El cargador ocupa todo el ancho disponible */
    [data-testid="stFileUploader"] {
        width: 100%;
    }

    /* Evita que tablas anchas rompan el diseño */
    [data-testid="stDataFrame"] {
        max-width: 100%;
        overflow-x: auto;
    }

    /* CELULAR / TABLET PEQUEÑA */
    @media (max-width: 768px) {
        /* Reserva espacio real debajo de la barra superior fija de Streamlit.
           Esto evita que el título quede oculto al abrir la app normalmente. */
        [data-testid="stAppViewBlockContainer"],
        .block-container {
            padding-top: 4.25rem !important;
            padding-left: 0.75rem !important;
            padding-right: 0.75rem !important;
            padding-bottom: 1.5rem !important;
            max-width: 100% !important;
        }

        h1 {
            font-size: 1.55rem !important;
            line-height: 1.28 !important;
            margin-top: 0 !important;
            padding-top: 0 !important;
            overflow: visible !important;
        }

        h2, h3 {
            line-height: 1.25 !important;
        }

        /* Convierte las filas de columnas de Streamlit en bloques verticales */
        div[data-testid="stHorizontalBlock"] {
            flex-direction: column !important;
            gap: 0.65rem !important;
        }

        div[data-testid="column"] {
            width: 100% !important;
            flex: 1 1 100% !important;
            min-width: 100% !important;
        }

        [data-testid="stMetric"] {
            width: 100%;
            min-height: auto;
            padding: 0.75rem 0.9rem;
        }

        [data-testid="stMetricValue"] {
            font-size: 1.45rem !important;
        }

        .stDownloadButton,
        .stDownloadButton > button,
        .stButton,
        .stButton > button {
            width: 100% !important;
        }

        [data-testid="stFileUploaderDropzone"] {
            padding: 0.85rem !important;
        }

        /* Mantiene la tabla consultable con desplazamiento horizontal */
        [data-testid="stDataFrame"] {
            overflow-x: auto !important;
            -webkit-overflow-scrolling: touch;
        }

        /* Reduce espacios excesivos en mensajes */
        [data-testid="stAlert"] {
            padding: 0.75rem !important;
        }
    }
    </style>
    """,
    unsafe_allow_html=True,
)


st.title("🏛️ Seguimiento de Inversiones usando datos del SIAF")

st.caption("Municipalidad Provincial de San Román - Juliaca")



# ---------------------------------------------------------

# FUNCIONES DE PROCESAMIENTO Y LÓGICA DE DATOS

# ---------------------------------------------------------

def extraer_cui(texto):

    if pd.isna(texto):

        return ""

    coincidencia = re.search(r"\b\d{7}\b", str(texto))

    if coincidencia:

        return coincidencia.group(0)

    coincidencia_inicio = re.search(r"^\d+", str(texto))

    return coincidencia_inicio.group(0) if coincidencia_inicio else str(texto)





def limpiar_nombre_proyecto(texto):

    if pd.isna(texto):

        return ""

    return re.sub(r"^\d+\\.\s*", "", str(texto)).strip()





def determinar_tipo(cui):

    cui_str = str(cui).strip()

    if cui_str.startswith("2"):

        return "PROYECTO"

    elif cui_str.startswith("3"):

        return "ACTIVIDAD"

    elif len(cui_str) == 7 and cui_str.isdigit():

        return "PROYECTO"

    else:

        return "ACTIVIDAD"





def obtener_barra_app(porcentaje):

    try:

        pct = float(porcentaje)

    except (ValueError, TypeError):

        pct = 0.0



    pct = min(max(pct, 0.0), 100.0)



    if pct < 30.0:

        icono = "🟥"

    elif pct <= 50.0:

        icono = "🟧"

    else:

        icono = "🟩"



    num_bloques = int(round((pct / 100.0) * 10))

    if pct > 0 and num_bloques == 0:

        num_bloques = 1



    return (icono * num_bloques) + ("⬛" * (10 - num_bloques))





# ---------------------------------------------------------

# FUNCIÓN PARA GENERAR PDF DE REPORTE

# ---------------------------------------------------------

def generar_pdf_reporte(df_resumen_pdf, datos_kpi, texto_interpretacion, fecha_consulta):

    buffer = io.BytesIO()

    doc = SimpleDocTemplate(

        buffer,

        pagesize=landscape(A4),

        rightMargin=20,

        leftMargin=20,

        topMargin=20,

        bottomMargin=20,

    )

    elements = []

    styles = getSampleStyleSheet()



    # Estilos de texto

    titulo_style = ParagraphStyle(

        "TituloPDF",

        parent=styles["Heading1"],

        fontSize=16,

        leading=20,

        textColor=colors.HexColor("#1A2B4C"),

        alignment=1,  # Centrado

    )

    subtitulo_style = ParagraphStyle(

        "SubtituloPDF",

        parent=styles["Normal"],

        fontSize=10,

        leading=13,

        textColor=colors.HexColor("#555555"),

        alignment=1,

    )

    seccion_style = ParagraphStyle(

        "SeccionPDF",

        parent=styles["Heading2"],

        fontSize=12,

        leading=15,

        textColor=colors.HexColor("#1A2B4C"),

        spaceBefore=10,

        spaceAfter=5,

    )

    body_style = ParagraphStyle(

        "BodyPDF", parent=styles["Normal"], fontSize=8, leading=11

    )

    header_tabla_style = ParagraphStyle(

        "HeaderTablaPDF",

        parent=styles["Normal"],

        fontSize=7,

        leading=9,

        fontName="Helvetica-Bold",

        textColor=colors.white,

        alignment=1,

    )

    cell_tabla_style = ParagraphStyle(

        "CellTablaPDF", parent=styles["Normal"], fontSize=6.5, leading=8

    )



    # 1. Encabezado

    elements.append(

        Paragraph("REPORTE DE SEGUIMIENTO PRESUPUESTAL SIAF", titulo_style)

    )

    elements.append(

        Paragraph(

            "Municipalidad Provincial de San Román - Juliaca", subtitulo_style

        )

    )

    elements.append(Paragraph(f"Fecha de consulta: {fecha_consulta}", subtitulo_style))

    elements.append(Spacer(1, 10))



    # 2. Resumen KPI

    elements.append(Paragraph("1. Resumen Presupuestal General", seccion_style))

    tabla_kpi_data = [

        [

            Paragraph("**PIM Total:**", body_style),

            Paragraph(f"S/ {datos_kpi['pim']:,.2f}", body_style),

            Paragraph("**Saldo por Certificar:**", body_style),

            Paragraph(f"S/ {datos_kpi['saldo_cert']:,.2f}", body_style),

        ],

        [

            Paragraph("**Certificado:**", body_style),

            Paragraph(f"S/ {datos_kpi['cert']:,.2f}", body_style),

            Paragraph("**Saldo por Devengar:**", body_style),

            Paragraph(f"S/ {datos_kpi['saldo_dev']:,.2f}", body_style),

        ],

        [

            Paragraph("**Compromiso Anual:**", body_style),

            Paragraph(f"S/ {datos_kpi['comp']:,.2f}", body_style),

            Paragraph("**Avance Global (%):**", body_style),

            Paragraph(f"{datos_kpi['pct_global']:.2f}%", body_style),

        ],

        [

            Paragraph("**Devengado:**", body_style),

            Paragraph(f"S/ {datos_kpi['dev']:,.2f}", body_style),

            Paragraph("", body_style),

            Paragraph("", body_style),

        ],

    ]

    t_kpi = Table(tabla_kpi_data, colWidths=[120, 150, 120, 150])

    t_kpi.setStyle(

        TableStyle(

            [

                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F4F6F9")),

                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),

                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),

                ("PADDING", (0, 0), (-1, -1), 4),

            ]

        )

    )

    elements.append(t_kpi)

    elements.append(Spacer(1, 10))



    # 3. Interpretación Escrita

    elements.append(Paragraph("2. Análisis de Ejecución", seccion_style))

    texto_limpio = texto_interpretacion.replace("*", "•")

    for linea in texto_limpio.split("\n"):

        linea_str = linea.strip()

        if linea_str:

            elements.append(Paragraph(linea_str, body_style))

            elements.append(Spacer(1, 2))

    elements.append(Spacer(1, 10))



    # 4. Tabla de Detalle

    elements.append(

        Paragraph("3. Detalle por Proyecto y Actividad", seccion_style)

    )



    headers = [

        "CUI",

        "PROYECTO / ACTIVIDAD",

        "TIPO",

        "PIM (S/)",

        "CERT. (S/)",

        "COMP. (S/)",

        "DEV. (S/)",

        "SALDO CERT. (S/)",

        "SALDO DEV. (S/)",

        "AVANCE %",

    ]

    data_tabla = [[Paragraph(h, header_tabla_style) for h in headers]]



    for idx, row in df_resumen_pdf.iterrows():

        is_total = str(row["CUI"]) == "TOTAL"

        estilo_celda = (

            ParagraphStyle(

                "CellTotal",

                parent=cell_tabla_style,

                fontName="Helvetica-Bold",

            )

            if is_total

            else cell_tabla_style

        )



        f_row = [

            Paragraph(str(row["CUI"]), estilo_celda),

            Paragraph(str(row["PROYECTO / ACTIVIDAD"]), estilo_celda),

            Paragraph(str(row["TIPO"]), estilo_celda),

            Paragraph(

                f"{row[df_resumen_pdf.columns[3]]:,.2f}", estilo_celda

            ),

            Paragraph(

                f"{row[df_resumen_pdf.columns[4]]:,.2f}", estilo_celda

            ),

            Paragraph(

                f"{row[df_resumen_pdf.columns[5]]:,.2f}", estilo_celda

            ),

            Paragraph(

                f"{row[df_resumen_pdf.columns[6]]:,.2f}", estilo_celda

            ),

            Paragraph(

                f"{row['SALDO_POR_CERTIFICAR']:,.2f}", estilo_celda

            ),

            Paragraph(

                f"{row['SALDO_POR_DEVENGAR']:,.2f}", estilo_celda

            ),

            Paragraph(f"{row['PORCENTAJE_AVANCE']:.2f}%", estilo_celda),

        ]

        data_tabla.append(f_row)



    col_widths = [50, 180, 55, 60, 60, 60, 60, 65, 65, 50]

    t_detalle = Table(data_tabla, colWidths=col_widths, repeatRows=1)



    t_style = [

        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A2B4C")),

        ("ALIGN", (0, 0), (-1, -1), "LEFT"),

        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),

        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CCCCCC")),

        ("TOPPADDING", (0, 0), (-1, -1), 2),

        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),

    ]



    if len(data_tabla) > 1:

        t_style.append(

            (

                "BACKGROUND",

                (0, -1),

                (-1, -1),

                colors.HexColor("#E5E7EB"),

            )

        )



    t_detalle.setStyle(TableStyle(t_style))

    elements.append(t_detalle)



    doc.build(elements)

    buffer.seek(0)

    return buffer.getvalue()





# ---------------------------------------------------------

# CARGA DE ARCHIVOS Y PROCESAMIENTO COMPLETO

# ---------------------------------------------------------

archivo_data2 = st.file_uploader(

    "📁 Cargar Reporte de Consulta Amigable / SIAF (.xlsx, .xls)", type=["xlsx", "xls"]

)



if archivo_data2:

    try:

        # Fecha y hora exactas en que se procesa la consulta (hora de Perú)
        fecha_consulta_dt = datetime.now(ZoneInfo("America/Lima"))
        fecha_consulta = fecha_consulta_dt.strftime("%d/%m/%Y - %H:%M:%S")
        fecha_archivo = fecha_consulta_dt.strftime("%d-%m-%Y_%H-%M-%S")

        df = pd.read_excel(archivo_data2)

        df.columns = df.columns.astype(str).str.strip().str.upper()



        # Identificar CUI

        col_cui = [

            c for c in df.columns if "PRODUCTO" in c or "PROYECTO" in c or "CUI" in c

        ]

        nombre_col_cui = col_cui[0] if col_cui else df.columns[0]



        df["CUI"] = df[nombre_col_cui].apply(extraer_cui)

        df["PROYECTO / ACTIVIDAD"] = df[nombre_col_cui].apply(limpiar_nombre_proyecto)



        # Identificar PIM

        col_pim = [c for c in df.columns if "PIM" in c or "MTO_PIM" in c]

        nombre_pim = col_pim[0] if col_pim else "PIM"

        df[nombre_pim] = (

            pd.to_numeric(df[nombre_pim], errors="coerce").fillna(0)

            if nombre_pim in df.columns

            else 0.0

        )



        # Identificar Certificado

        col_cert = [

            c for c in df.columns if "CERTIFICADO" in c or "MTO_CERTIFICADO" in c or "CERT" in c

        ]

        nombre_cert = "CERTIFICADO"

        if col_cert:

            for c in col_cert:

                df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)

            df[nombre_cert] = df[col_cert].sum(axis=1)

        else:

            df[nombre_cert] = 0.0



        # Identificar Compromiso Anual

        col_comp = [

            c for c in df.columns if "COMPRO" in c or "MTO_COMPRO_ANUAL" in c or "COMPROMISO" in c

        ]

        nombre_comp = "COMPROMISO ANUAL"

        if col_comp:

            for c in col_comp:

                df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)

            df[nombre_comp] = df[col_comp].sum(axis=1)

        else:

            df[nombre_comp] = 0.0



        # Identificar Devengado

        cols_devengado = [c for c in df.columns if "DEVENGADO" in c or "DEV" in c]

        if cols_devengado:

            for c in cols_devengado:

                df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)

            df["DEVENGADO"] = df[cols_devengado].sum(axis=1)

        else:

            cols_excluir = [

                nombre_col_cui,

                "CUI",

                "PROYECTO / ACTIVIDAD",

                nombre_pim,

                nombre_cert,

                nombre_comp,

            ]

            cols_numericas = [c for c in df.columns if c not in cols_excluir]

            for c in cols_numericas:

                df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)

            df["DEVENGADO"] = df[cols_numericas].sum(axis=1)



        nombre_dev = "DEVENGADO"



        df["TIPO"] = df["CUI"].apply(determinar_tipo)



        # Agrupar

        df_agrupado = (

            df.groupby(["CUI", "PROYECTO / ACTIVIDAD", "TIPO"])[

                [nombre_pim, nombre_cert, nombre_comp, nombre_dev]

            ]

            .sum()

            .reset_index()

        )



        # Cálculos de Saldos

        df_agrupado["SALDO_POR_CERTIFICAR"] = df_agrupado[nombre_pim] - df_agrupado[nombre_cert]

        df_agrupado["SALDO_POR_DEVENGAR"] = df_agrupado[nombre_pim] - df_agrupado[nombre_dev]

        df_agrupado["PORCENTAJE_AVANCE"] = (

            (df_agrupado[nombre_dev] / df_agrupado[nombre_pim]) * 100

        ).round(2).fillna(0)

        df_agrupado["Gráfico Avance"] = df_agrupado["PORCENTAJE_AVANCE"].apply(obtener_barra_app)



        st.success("✅ Archivo procesado con éxito.")
        st.caption(f"🗓️ Fecha de consulta: {fecha_consulta}")



        # Totales

        total_pim = df_agrupado[nombre_pim].sum()

        total_cert = df_agrupado[nombre_cert].sum()

        total_comp = df_agrupado[nombre_comp].sum()

        total_dev = df_agrupado[nombre_dev].sum()

        total_saldo_cert = df_agrupado["SALDO_POR_CERTIFICAR"].sum()

        total_saldo_dev = df_agrupado["SALDO_POR_DEVENGAR"].sum()

        pct_global = (total_dev / total_pim * 100) if total_pim > 0 else 0.0

        pct_cert_global = (total_cert / total_pim * 100) if total_pim > 0 else 0.0



        # Totales por Tipo

        proyectos_df = df_agrupado[df_agrupado["TIPO"] == "PROYECTO"]

        num_proyectos = len(proyectos_df)

        dev_proyectos = proyectos_df[nombre_dev].sum()

        pim_proyectos = proyectos_df[nombre_pim].sum()

        pct_proyectos = (dev_proyectos / pim_proyectos * 100) if pim_proyectos > 0 else 0.0



        # Tarjetas KPI

        st.subheader("📌 Resumen Presupuestal General")

        kpi1, kpi2, kpi3, kpi4 = st.columns(4)

        kpi1.metric("PIM Total", f"S/ {total_pim:,.2f}")

        kpi2.metric("Certificado Total", f"S/ {total_cert:,.2f}")

        kpi3.metric("Compromiso Anual", f"S/ {total_comp:,.2f}")

        kpi4.metric("Devengado Total", f"S/ {total_dev:,.2f}")



        kpi5, kpi6, kpi7 = st.columns(3)

        kpi5.metric("Saldo Por Certificar", f"S/ {total_saldo_cert:,.2f}")

        kpi6.metric("Saldo Por Devengar", f"S/ {total_saldo_dev:,.2f}")

        kpi7.metric(

            "Avance Global %",

            f"{pct_global:.2f}%",

            delta=f"{pct_global:.2f}%",

            delta_color="normal" if pct_global >= 50 else "inverse",

        )



        # ---------------------------------------------------------

        # SECCIÓN DE INTERPRETACIÓN EJECUTIVA

        # ---------------------------------------------------------

        st.markdown("---")

        st.subheader("📝 Estado Actual y Análisis de Ejecución")



        if pct_global < 30:

            nivel_estado = "CRÍTICO (Bajo avance)"

        elif pct_global <= 50:

            nivel_estado = "REGULAR (Avance en proceso)"

        else:

            nivel_estado = "ACEPTABLE / OPTIMO"



        texto_interpretacion = f"""

        * **Nivel de Ejecución:** El avance presupuestal global de la entidad se encuentra en un estado **{nivel_estado}** con un **{pct_global:.2f}%** devengado respecto al PIM Total.

        * **Monto Ejecutado:** De un Presupuesto Institucional Modificado (PIM) de **S/ {total_pim:,.2f}**, se ha devengado a la fecha un monto total de **S/ {total_dev:,.2f}**, dejando un saldo pendiente por devengar de **S/ {total_saldo_dev:,.2f}**.

        * **Certificación Presupuestal:** Se ha logrado certificar el **{pct_cert_global:.2f}%** del presupuesto disponible (**S/ {total_cert:,.2f}**), existiendo aún un saldo sin certificar de **S/ {total_saldo_cert:,.2f}**.

        * **Inversión en Proyectos:** En la cartera de proyectos de inversión, se registra un total de **{num_proyectos} registros**, con un avance acumulado del **{pct_proyectos:.2f}%** (Devengado: S/ {dev_proyectos:,.2f} de un PIM de S/ {pim_proyectos:,.2f}).

        """

        st.info(texto_interpretacion)



        st.divider()



        # Reordenar columnas

        columnas_ordenadas = [

            "CUI",

            "PROYECTO / ACTIVIDAD",

            "TIPO",

            nombre_pim,

            nombre_cert,

            nombre_comp,

            nombre_dev,

            "SALDO_POR_CERTIFICAR",

            "SALDO_POR_DEVENGAR",

            "PORCENTAJE_AVANCE",

            "Gráfico Avance",

        ]

        df_agrupado = df_agrupado[columnas_ordenadas]



        # Vista preliminar

        df_vista = df_agrupado.copy()

        fila_total = {col: "" for col in df_vista.columns}

        fila_total["CUI"] = "TOTAL"

        fila_total["PROYECTO / ACTIVIDAD"] = "TOTAL GENERAL"

        fila_total[nombre_pim] = total_pim

        fila_total[nombre_cert] = total_cert

        fila_total[nombre_comp] = total_comp

        fila_total[nombre_dev] = total_dev

        fila_total["SALDO_POR_CERTIFICAR"] = total_saldo_cert

        fila_total["SALDO_POR_DEVENGAR"] = total_saldo_dev

        fila_total["PORCENTAJE_AVANCE"] = round(pct_global, 2)

        fila_total["Gráfico Avance"] = obtener_barra_app(pct_global)



        df_vista = pd.concat([df_vista, pd.DataFrame([fila_total])], ignore_index=True)



        # ---------------------------------------------------------

        # GENERACIÓN DE ARCHIVOS PARA DESCARGA

        # ---------------------------------------------------------

        # 1. Excel

        df_excel = df_agrupado.copy()

        df_excel.drop(columns=["Gráfico Avance"], inplace=True, errors="ignore")

        # Evita que la fila TOTAL vuelva a crear una columna vacía "Gráfico Avance"
        # al concatenarse con df_excel. Solo se conservan las columnas ya existentes.
        fila_total_excel = {col: fila_total.get(col, "") for col in df_excel.columns}

        df_excel = pd.concat(
            [df_excel, pd.DataFrame([fila_total_excel])], ignore_index=True
        )



        buffer_excel = io.BytesIO()

        with pd.ExcelWriter(buffer_excel, engine="openpyxl") as writer:

            df_excel.to_excel(writer, index=False, sheet_name="Seguimiento")

            ws = writer.sheets["Seguimiento"]



            # Hoja informativa adicional: no altera la estructura de "Seguimiento"
            ws_info = writer.book.create_sheet("Información")
            ws_info["A1"] = "FECHA DE CONSULTA"
            ws_info["B1"] = fecha_consulta
            ws_info["A1"].font = Font(bold=True)
            ws_info.column_dimensions["A"].width = 24
            ws_info.column_dimensions["B"].width = 24

            headers = [cell.value for cell in ws[1]]

            formato_moneda = '"S/" #,##0.00'

            formato_pct = "0.00%"

            col_pct_idx = None



            for col_idx, col_name in enumerate(headers, 1):

                if col_name in [

                    nombre_pim,

                    nombre_cert,

                    nombre_comp,

                    nombre_dev,

                    "SALDO_POR_CERTIFICAR",

                    "SALDO_POR_DEVENGAR",

                ]:

                    for row_idx in range(2, ws.max_row + 1):

                        cell = ws.cell(row=row_idx, column=col_idx)

                        cell.number_format = formato_moneda

                elif col_name == "PORCENTAJE_AVANCE":

                    col_pct_idx = col_idx

                    for row_idx in range(2, ws.max_row + 1):

                        cell = ws.cell(row=row_idx, column=col_idx)

                        cell.value = (cell.value or 0) / 100

                        cell.number_format = formato_pct



            ultima_fila = ws.max_row

            for col_idx in range(1, len(headers) + 1):

                cell_tot = ws.cell(row=ultima_fila, column=col_idx)

                cell_tot.font = Font(bold=True)



            col_grafico_idx = len(headers) + 1

            col_pct_letter = openpyxl.utils.get_column_letter(col_pct_idx)

            col_grafico_letter = openpyxl.utils.get_column_letter(col_grafico_idx)



            ws.cell(row=1, column=col_grafico_idx, value="GRAFICO_AVANCE")



            for row_idx in range(2, ws.max_row + 1):

                cell_graf = ws.cell(row=row_idx, column=col_grafico_idx)

                cell_graf.value = f'=REPT("█", ROUND({col_pct_letter}{row_idx}*10, 0))'



            font_rojo = Font(color="C0392B", bold=True)

            font_naranja = Font(color="D35400", bold=True)

            font_verde = Font(color="27AE60", bold=True)



            fill_rojo = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")

            fill_naranja = PatternFill(start_color="FCF3CF", end_color="FCF3CF", fill_type="solid")

            fill_verde = PatternFill(start_color="D4EFDF", end_color="D4EFDF", fill_type="solid")



            rango_pct = f"{col_pct_letter}2:{col_pct_letter}{ws.max_row}"

            rango_grafico = f"{col_grafico_letter}2:{col_grafico_letter}{ws.max_row}"



            ws.conditional_formatting.add(

                rango_pct,

                CellIsRule(operator="lessThan", formula=["0.30"], stopIfTrue=True, fill=fill_rojo),

            )

            ws.conditional_formatting.add(

                rango_pct,

                CellIsRule(

                    operator="between", formula=["0.30", "0.50"], stopIfTrue=True, fill=fill_naranja

                ),

            )

            ws.conditional_formatting.add(

                rango_pct,

                CellIsRule(

                    operator="greaterThan", formula=["0.50"], stopIfTrue=True, fill=fill_verde

                ),

            )



            ws.conditional_formatting.add(

                rango_grafico,

                CellIsRule(

                    operator="lessThan", formula=["0.30"], stopIfTrue=True, font=font_rojo

                ),

            )

            ws.conditional_formatting.add(

                rango_grafico,

                CellIsRule(

                    operator="between", formula=["0.30", "0.50"], stopIfTrue=True, font=font_naranja

                ),

            )

            ws.conditional_formatting.add(

                rango_grafico,

                CellIsRule(

                    operator="greaterThan", formula=["0.50"], stopIfTrue=True, font=font_verde

                ),

            )



            # Mejora visual de la columna PROYECTO / ACTIVIDAD:
            # mantiene el nombre completo, limita el ancho visible y permite
            # consultar el texto completo mediante una nota de Excel.
            col_proyecto_idx = headers.index("PROYECTO / ACTIVIDAD") + 1
            col_proyecto_letter = openpyxl.utils.get_column_letter(col_proyecto_idx)

            for row_idx in range(2, ws.max_row + 1):
                cell_proyecto = ws.cell(row=row_idx, column=col_proyecto_idx)
                texto_proyecto = str(cell_proyecto.value or "")
                cell_proyecto.alignment = Alignment(
                    wrap_text=True, vertical="top"
                )

                # Si el nombre es largo, se conserva íntegro en una nota.
                # Al pasar el cursor o seleccionar la celda puede revisarse completo.
                if len(texto_proyecto) > 60:
                    cell_proyecto.comment = Comment(
                        texto_proyecto,
                        "Seguimiento SIAF",
                    )

                # Altura compacta para no ocupar demasiado espacio en pantalla.
                if row_idx < ws.max_row:
                    ws.row_dimensions[row_idx].height = 32

            for col in ws.columns:

                max_len = max(len(str(cell.value or "")) for cell in col)

                col_letter = openpyxl.utils.get_column_letter(col[0].column)

                # Evita que los nombres largos de proyecto ensanchen toda la hoja.
                if col_letter == col_proyecto_letter:
                    ws.column_dimensions[col_letter].width = 45
                else:
                    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)



        # 2. PDF

        dict_kpis = {

            "pim": total_pim,

            "cert": total_cert,

            "comp": total_comp,

            "dev": total_dev,

            "saldo_cert": total_saldo_cert,

            "saldo_dev": total_saldo_dev,

            "pct_global": pct_global,

        }

        df_pdf = df_excel.copy()

        bytes_pdf = generar_pdf_reporte(df_pdf, dict_kpis, texto_interpretacion, fecha_consulta)



        # Botones de descarga alineados

        c_down1, c_down2 = st.columns(2)



        with c_down1:

            st.download_button(

                label="📥 Descargar Reporte Excel (.xlsx)",

                data=buffer_excel.getvalue(),

                file_name=f"Seguimiento_Procesado_San_Roman_{fecha_archivo}.xlsx",

                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",

                type="primary",

                use_container_width=True,

            )



        with c_down2:

            st.download_button(

                label="📄 Descargar Reporte PDF (.pdf)",

                data=bytes_pdf,

                file_name=f"Reporte_Seguimiento_San_Roman_{fecha_archivo}.pdf",

                mime="application/pdf",

                type="primary",

                use_container_width=True,

            )



        st.subheader("📋 Vista previa de la hoja Seguimiento")

        st.dataframe(

            df_vista,

            column_config={

                "CUI": st.column_config.TextColumn("CUI"),

                "PROYECTO / ACTIVIDAD": st.column_config.TextColumn("Proyecto / Actividad"),

                "TIPO": st.column_config.TextColumn("Tipo"),

                nombre_pim: st.column_config.NumberColumn("PIM (S/)", format="S/ %,.2f"),

                nombre_cert: st.column_config.NumberColumn("Certificado (S/)", format="S/ %,.2f"),

                nombre_comp: st.column_config.NumberColumn(

                    "Compromiso Anual (S/)", format="S/ %,.2f"

                ),

                nombre_dev: st.column_config.NumberColumn("Devengado (S/)", format="S/ %,.2f"),

                "SALDO_POR_CERTIFICAR": st.column_config.NumberColumn(

                    "Saldo por Certificar (S/)", format="S/ %,.2f"

                ),

                "SALDO_POR_DEVENGAR": st.column_config.NumberColumn(

                    "Saldo por Devengar (S/)", format="S/ %,.2f"

                ),

                "PORCENTAJE_AVANCE": st.column_config.NumberColumn(

                    "AVANCE %", format="%.2f %%"

                ),

                "Gráfico Avance": st.column_config.TextColumn("Gráfico Avance"),

            },

            use_container_width=True,

        )



    except Exception as e:

        st.error(f"Error procesando el archivo: {e}")

else:

    st.info("Por favor sube el archivo de Información Presupuestaria del gasto del SIAF para iniciar.")